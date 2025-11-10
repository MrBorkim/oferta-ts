#!/usr/bin/env python3
"""
Serwer FastAPI do generowania ofert z szablonów DOCX.
Silnik szablonów: docxtpl (Jinja2)
Konwersja: DOCX → PDF (LibreOffice) → JPG (PyMuPDF, 100 dpi)
"""

import os
import json
import subprocess
import tempfile
import shutil
from pathlib import Path
from typing import Optional, List, Dict, Any
from io import BytesIO
import zipfile

from fastapi import FastAPI, HTTPException, Header, Request
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel, Field
import fitz  # PyMuPDF
from PIL import Image
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm

# ========== KONFIGURACJA Z ENV ==========
API_KEY = os.getenv("API_KEY", "devkey")
PORT = int(os.getenv("PORT", "7077"))
HOST = os.getenv("HOST", "0.0.0.0")
TEMPLATES_ROOT = Path(os.getenv("TEMPLATES_ROOT", "templates"))
PRODUCTS_ROOT = Path(os.getenv("PRODUCTS_ROOT", "products"))
JPEG_QUALITY = int(os.getenv("JPEG_QUALITY", "85"))
DPI = int(os.getenv("DPI", "100"))

# ========== FASTAPI APP ==========
app = FastAPI(
    title="Offer Rendering API",
    description="Generowanie ofert z szablonów DOCX → PDF → JPG",
    version="1.0.0"
)


# ========== MODELE PYDANTIC ==========
class ProductItem(BaseModel):
    product_id: str = Field(..., description="ID produktu (nazwa katalogu w PRODUCTS_ROOT)")
    page: Optional[int] = Field(1, description="Numer strony, gdzie ma się pojawić produkt")
    slot: Optional[str] = Field("main", description="Slot pozycji na stronie (hero, main, grid_1, etc.)")
    sequence: Optional[int] = Field(None, description="Kolejność wyświetlania")
    image: Optional[str] = Field("cover.jpg", description="Nazwa pliku obrazu (domyślnie cover.jpg)")
    data: Optional[Dict[str, Any]] = Field(default_factory=dict, description="Dodatkowe dane dla produktu")


class RenderRequest(BaseModel):
    template: str = Field(..., description="Nazwa szablonu (folder w TEMPLATES_ROOT)")
    placeholders: Dict[str, Any] = Field(default_factory=dict, description="Placeholders do podstawienia w szablonie")
    products: List[ProductItem] = Field(default_factory=list, description="Lista produktów do wstawienia")
    return_mode: Optional[str] = Field("zip", description="'first_page_inline' lub 'zip' (domyślnie)")


# ========== MIDDLEWARE: BEZPIECZEŃSTWO ==========
@app.middleware("http")
async def verify_api_key(request: Request, call_next):
    """Weryfikacja X-API-Key dla wszystkich endpointów poza /health"""
    if request.url.path == "/health":
        return await call_next(request)

    api_key = request.headers.get("X-API-Key")
    if api_key != API_KEY:
        return Response(
            content=json.dumps({"detail": "Invalid or missing X-API-Key"}),
            status_code=401,
            media_type="application/json"
        )

    return await call_next(request)


# ========== ENDPOINT: HEALTH ==========
@app.get("/health")
def health_check():
    """Sprawdzenie czy serwer działa i czy LibreOffice jest dostępny"""
    libreoffice_available = check_libreoffice()

    return {
        "status": "ok",
        "libreoffice_available": libreoffice_available,
        "templates_root": str(TEMPLATES_ROOT.absolute()),
        "products_root": str(PRODUCTS_ROOT.absolute()),
        "dpi": DPI,
        "jpeg_quality": JPEG_QUALITY
    }


# ========== ENDPOINT: RENDER ==========
@app.post("/render")
def render_offer(req: RenderRequest):
    """
    Główny endpoint do renderowania ofert.

    1. Renderuje szablon DOCX z docxtpl
    2. Konwertuje DOCX → PDF (LibreOffice)
    3. Konwertuje PDF → JPG (PyMuPDF, 100 dpi)
    4. Zwraca pierwszą stronę jako image/jpeg lub ZIP z wszystkimi stronami
    """

    # Walidacja: czy szablon istnieje
    template_path = TEMPLATES_ROOT / req.template
    if not template_path.exists() or not template_path.is_dir():
        raise HTTPException(
            status_code=404,
            detail=f"Template folder not found: {req.template}. Check TEMPLATES_ROOT={TEMPLATES_ROOT}"
        )

    # Sprawdź czy LibreOffice jest dostępny
    if not check_libreoffice():
        raise HTTPException(
            status_code=500,
            detail="LibreOffice not found. Install: brew install libreoffice (macOS) or apt-get install libreoffice (Linux)"
        )

    # Tymczasowe katalogi/pliki
    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir_path = Path(tmpdir)

        try:
            # 1. Przygotuj kontekst dla docxtpl
            context = prepare_context(req.placeholders, req.products)

            # 2. Renderuj szablon(y) DOCX
            rendered_docx = render_template(template_path, context, tmpdir_path)

            # 3. Konwertuj DOCX → PDF
            pdf_path = convert_docx_to_pdf(rendered_docx, tmpdir_path)

            # 4. Konwertuj PDF → JPG
            jpg_paths = convert_pdf_to_jpg(pdf_path, tmpdir_path, dpi=DPI, quality=JPEG_QUALITY)

            # 5. Zwróć wynik
            if req.return_mode == "first_page_inline":
                # Zwróć tylko pierwszą stronę jako image/jpeg
                with open(jpg_paths[0], "rb") as f:
                    image_data = f.read()
                return Response(content=image_data, media_type="image/jpeg")
            else:
                # Zwróć wszystkie strony jako ZIP
                zip_buffer = create_zip(jpg_paths)
                return StreamingResponse(
                    BytesIO(zip_buffer),
                    media_type="application/zip",
                    headers={"Content-Disposition": f"attachment; filename=offer_{req.template}.zip"}
                )

        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Error rendering offer: {str(e)}")


# ========== FUNKCJE POMOCNICZE ==========

def check_libreoffice() -> bool:
    """Sprawdza czy LibreOffice jest dostępny w systemie"""
    try:
        result = subprocess.run(
            ["libreoffice", "--version"],
            capture_output=True,
            text=True,
            timeout=5
        )
        return result.returncode == 0
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return False


def prepare_context(placeholders: Dict[str, Any], products: List[ProductItem]) -> Dict[str, Any]:
    """
    Przygotowuje kontekst dla docxtpl.

    Returns:
        {
            "data": placeholders,
            "products": lista produktów z dodanym image_abs
        }
    """
    # Przygotuj listę produktów z absolutnymi ścieżkami do obrazów
    products_list = []
    for product in products:
        product_dict = product.dict()

        # Znajdź obraz produktu
        product_dir = PRODUCTS_ROOT / product.product_id
        if not product_dir.exists():
            raise ValueError(f"Product directory not found: {product.product_id}")

        # Dodaj absolutną ścieżkę do obrazu
        image_name = product.image or "cover.jpg"
        image_path = product_dir / image_name

        if image_path.exists():
            product_dict["image_abs"] = str(image_path.absolute())
        else:
            # Jeśli obraz nie istnieje, można pominąć lub rzucić błąd
            product_dict["image_abs"] = None

        products_list.append(product_dict)

    return {
        "data": placeholders,
        "products": products_list
    }


def render_template(template_path: Path, context: Dict[str, Any], tmpdir: Path) -> Path:
    """
    Renderuje szablon DOCX używając docxtpl.

    Args:
        template_path: Ścieżka do folderu szablonu
        context: Kontekst dla docxtpl (data + products)
        tmpdir: Katalog tymczasowy

    Returns:
        Ścieżka do wyrenderowanego pliku DOCX
    """

    # Sprawdź czy to szablon jednoplikowy czy wieloplikowy
    config_files = list(template_path.glob("*.json"))

    # Znajdź główny plik DOCX
    # Priorytet: oferta1.docx, wolftax.docx, lub pierwszy *.docx
    docx_files = list(template_path.glob("*.docx"))
    if not docx_files:
        raise ValueError(f"No DOCX files found in template: {template_path}")

    # Wybierz główny plik DOCX
    main_docx = None
    for candidate in ["oferta1.docx", "wolftax.docx", "Dok1.docx"]:
        candidate_path = template_path / candidate
        if candidate_path.exists():
            main_docx = candidate_path
            break

    if not main_docx:
        main_docx = docx_files[0]

    # Renderuj szablon
    doc = DocxTemplate(main_docx)

    # Dodaj funkcję InlineImage do kontekstu
    def create_inline_image(image_path: str, width_mm: int = 120):
        """Wrapper dla InlineImage"""
        if not image_path or not Path(image_path).exists():
            return ""
        return InlineImage(doc, image_path, width=Mm(width_mm))

    # Rozszerz kontekst o funkcję InlineImage
    context["InlineImage"] = create_inline_image

    # Renderuj
    try:
        doc.render(context)
    except Exception as e:
        raise ValueError(f"Error rendering template with docxtpl: {str(e)}")

    # Zapisz wyrenderowany DOCX
    output_path = tmpdir / "rendered.docx"
    doc.save(output_path)

    return output_path


def convert_docx_to_pdf(docx_path: Path, tmpdir: Path) -> Path:
    """
    Konwertuje DOCX → PDF używając LibreOffice headless.

    Args:
        docx_path: Ścieżka do pliku DOCX
        tmpdir: Katalog tymczasowy dla pliku PDF

    Returns:
        Ścieżka do pliku PDF
    """
    try:
        subprocess.run(
            [
                "libreoffice",
                "--headless",
                "--convert-to", "pdf",
                "--outdir", str(tmpdir),
                str(docx_path)
            ],
            check=True,
            capture_output=True,
            timeout=60
        )
    except subprocess.CalledProcessError as e:
        raise RuntimeError(f"LibreOffice conversion failed: {e.stderr.decode()}")
    except subprocess.TimeoutExpired:
        raise RuntimeError("LibreOffice conversion timed out (>60s)")

    # Znajdź wygenerowany PDF
    pdf_path = tmpdir / f"{docx_path.stem}.pdf"
    if not pdf_path.exists():
        raise RuntimeError(f"PDF not created by LibreOffice: {pdf_path}")

    return pdf_path


def convert_pdf_to_jpg(pdf_path: Path, tmpdir: Path, dpi: int = 100, quality: int = 85) -> List[Path]:
    """
    Konwertuje PDF → JPG używając PyMuPDF.

    Args:
        pdf_path: Ścieżka do pliku PDF
        tmpdir: Katalog tymczasowy dla plików JPG
        dpi: Rozdzielczość w DPI (domyślnie 100)
        quality: Jakość JPEG 0-100 (domyślnie 85)

    Returns:
        Lista ścieżek do plików JPG (page_001.jpg, page_002.jpg, ...)
    """
    jpg_paths = []

    try:
        pdf_doc = fitz.open(pdf_path)

        for page_num in range(len(pdf_doc)):
            page = pdf_doc[page_num]

            # Renderuj stronę do pixmap z odpowiednim DPI
            zoom = dpi / 72.0  # PyMuPDF używa 72 DPI jako bazę
            matrix = fitz.Matrix(zoom, zoom)
            pix = page.get_pixmap(matrix=matrix, alpha=False)

            # Konwertuj pixmap → PIL Image → JPEG z DPI
            img_data = pix.tobytes("jpeg", quality=quality)
            img = Image.open(BytesIO(img_data))

            # Zapisz z metadanymi DPI
            jpg_path = tmpdir / f"page_{page_num + 1:03d}.jpg"
            img.save(jpg_path, "JPEG", quality=quality, dpi=(dpi, dpi))

            jpg_paths.append(jpg_path)

        pdf_doc.close()

    except Exception as e:
        raise RuntimeError(f"Error converting PDF to JPG: {str(e)}")

    return jpg_paths


def create_zip(jpg_paths: List[Path]) -> bytes:
    """
    Tworzy archiwum ZIP z plików JPG.

    Args:
        jpg_paths: Lista ścieżek do plików JPG

    Returns:
        Zawartość pliku ZIP jako bytes
    """
    zip_buffer = BytesIO()

    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for jpg_path in jpg_paths:
            zip_file.write(jpg_path, arcname=jpg_path.name)

    return zip_buffer.getvalue()


# ========== MAIN ==========
if __name__ == "__main__":
    import uvicorn

    print(f"Starting Offer Rendering API on {HOST}:{PORT}")
    print(f"TEMPLATES_ROOT: {TEMPLATES_ROOT.absolute()}")
    print(f"PRODUCTS_ROOT: {PRODUCTS_ROOT.absolute()}")
    print(f"DPI: {DPI}, JPEG_QUALITY: {JPEG_QUALITY}")

    uvicorn.run(app, host=HOST, port=PORT)
